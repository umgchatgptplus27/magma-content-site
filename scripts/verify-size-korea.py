"""Recalculate publication aggregates from the locally supplied Size Korea workbook.
Requires the existing openpyxl environment. Does not export participant records.
Usage: python3 scripts/verify-size-korea.py
"""
import hashlib
import json
import statistics
from pathlib import Path

import openpyxl

root = Path(__file__).resolve().parents[1]
source = root / 'docs/size-korea/8차 인체치수조사(2020~24)_치수데이터(공개용).xlsx'
workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
sheet = workbook['(1~2차년도) 직접측정']
headers = next(sheet.iter_rows(min_row=7, max_row=7, values_only=True))
fields = {'neck_mm': '032. 목둘레', 'wrist_mm': '038. 손목둘레',
          'foot_length_mm': '122. 발직선길이', 'foot_width_mm': '123. 발너비',
          'height_mm': '002. 키', 'chest_mm': '041. 가슴둘레',
          'waist_mm': '045. 허리둘레', 'hip_mm': '048. 엉덩이둘레'}
columns = {key: next(i for i, h in enumerate(headers) if str(h).strip() == name)
           for key, name in fields.items()}
rows = [r for r in sheet.iter_rows(min_row=8, values_only=True)
        if r[3] == 'M' and isinstance(r[4], (int, float)) and 30 <= r[4] < 50]

def valid_measurement(value):
    return isinstance(value, (int, float)) and 0 < value < 9999

def summarize(values):
    values = sorted(values)
    def percentile(p):
        index = (len(values) - 1) * p
        lower = int(index)
        return values[lower] + (values[min(lower + 1, len(values) - 1)] - values[lower]) * (index - lower)
    return {'n': len(values), 'mean': statistics.mean(values),
            **{f'p{p}': percentile(p / 100) for p in [5, 25, 50, 75, 95]}}

result = {
    'verified_date': '2026-10-28', 'source_file': source.name,
    'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
    'source_url': 'https://sizekorea.kr/human-info/meas-report?measDegree=8',
    'sheet': sheet.title, 'filter': '성별 M, 30 <= 나이 < 50',
    'method': '비가중 집계. 항목별 숫자이며 0 < 값 < 9999인 값만 사용. 분위수는 (n-1)*p 위치의 선형 보간. 원자료 mm, cm 표시 시 10으로 나눔.',
    'limitations': '조사 참여자의 분포이며 한국 남성 전체의 가중 추정치나 개인의 핏·신발 폭 판정 기준이 아님. 원자료 파일의 공식 배포본 동일성은 별도 재확인 필요.',
    'paired_method': '같은 참여자의 두 유효 치수를 먼저 뺀 뒤 집계. 차이가 0 또는 음수여도 제외하지 않음. 집단별 분위수끼리 빼지 않음.',
    'groups': {}
}
for label, low, high in [('30-34', 30, 35), ('35-39', 35, 40), ('30-39', 30, 40), ('40-49', 40, 50), ('30-49', 30, 50)]:
    group = [r for r in rows if low <= r[4] < high]
    result['groups'][label] = {'participants': len(group), **{
        key: summarize([r[index] for r in group if valid_measurement(r[index])]) for key, index in columns.items()}}
    for key, left, right in [('chest_minus_waist_mm', 'chest_mm', 'waist_mm'),
                             ('hip_minus_waist_mm', 'hip_mm', 'waist_mm')]:
        result['groups'][label][key] = summarize([
            r[columns[left]] - r[columns[right]] for r in group
            if valid_measurement(r[columns[left]]) and valid_measurement(r[columns[right]])])
output = root / 'public/downloads/size-korea-fit-summary.json'
output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
print(json.dumps(result['groups'], ensure_ascii=False, indent=2))
